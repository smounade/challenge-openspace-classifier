from table.py import Table
import random
from openpyxl import Workbook

class Openspace:

    def __init__(self, tables: list, number_of_tables: int = 6):
        self.tables=tables
        self.number_of_tables=number_of_tables
  
    def organize(self, names: list):
        free_tables=[]
        for table in self.tables:
            if table.has_free_spot:
                free_tables.append(table)
        for name in names:
            table=random.choice(free_tables)
            table.assign_seat(name)
            
    def display(self):
        display=[]
        for table in self.tables:
            table_list=[]
            for seat in table:
                if seat.free:
                    table_list.append("free seat")
                else:
                    table_list.append(f"seat occupied by {seat.occupant}")
            display.append(table_list)
        print(display)
        return display

    def store(self, filename):
        # Create a workbook and select the active worksheet
        wb = Workbook()
        ws = wb.active

        # Write data row by row
        for row in self.display():
            ws.append(row)

        # Save the workbook to a file
        wb.save("output.xlsx")

        print("Data saved to output.xlsx")
